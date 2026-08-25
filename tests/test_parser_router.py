"""Tests for the parser_router format dispatcher + visual router.

DESIGN parser_router/DESIGN_parser_router.md (commit F6)
Each test exercises one branch of the dispatcher with a real in-memory
sample file. DeepDocClient is mocked so no container is required.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pytest

from agentic_rag_project.doc_processor.models import (
    BoundingBox,
    ParsedDoc,
)
from agentic_rag_project.doc_processor.parser import (
    DeepDocClient,
    parse_document_with_router,
)
from agentic_rag_project.doc_processor.text_extractors import (
    extract_docx,
    extract_pdf_text,
    extract_plain,
    extract_pptx,
    extract_xlsx,
    infer_format,
    is_scanned_pdf,
)
from agentic_rag_project.doc_processor.visual_router import DeepDocVisualRouter


# ---------------------------------------------------------------------------
# text_extractors — real in-process extraction on tiny sample files
# ---------------------------------------------------------------------------


def test_infer_format_accepts_all_supported_extensions() -> None:
    assert infer_format("a.pdf") == "pdf"
    assert infer_format("A.DOCX") == "docx"
    assert infer_format("b.pptx") == "pptx"
    assert infer_format("c.xlsx") == "xlsx"
    assert infer_format("readme.md") == "md"
    assert infer_format("log.txt") == "txt"


def test_infer_format_rejects_unknown_extension() -> None:
    with pytest.raises(ValueError):
        infer_format("malware.exe")


def test_extract_plain_markdown_splits_on_headings(tmp_path: Path) -> None:
    md = tmp_path / "doc.md"
    md.write_text(
        "# Top\nalpha\n\n## Sub\nbeta\n\n# Top2\ngamma\n",
        encoding="utf-8",
    )
    out = extract_plain(md)
    assert out.format == "md"
    # `## Sub` opens a new section (we split on every heading level).
    assert [s.heading for s in out.sections] == ["Top", "Sub", "Top2"]
    assert out.sections[0].level == 1
    assert out.sections[1].level == 2
    assert out.sections[2].blocks[0].text.strip() == "gamma"


def test_extract_plain_text_chunks_by_line_window(tmp_path: Path) -> None:
    txt = tmp_path / "doc.txt"
    txt.write_text("\n".join(f"line {i}" for i in range(75)), encoding="utf-8")
    out = extract_plain(txt)
    assert out.format == "txt"
    # 75 lines / 30 per section -> at least 2 sections
    assert len(out.sections) >= 2
    assert all(b.kind == "text" for s in out.sections for b in s.blocks)


def test_extract_xlsx_reads_sheet_with_headers(tmp_path: Path) -> None:
    from openpyxl import Workbook

    xlsx = tmp_path / "sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["name", "age"])
    ws.append(["alice", 30])
    ws.append(["bob", 25])
    wb.save(xlsx)
    wb.close()

    out = extract_xlsx(xlsx)
    assert out.format == "xlsx"
    assert len(out.sections) == 1
    assert out.sections[0].heading == "People"
    assert len(out.sections[0].blocks) == 2
    assert "name: alice" in out.sections[0].blocks[0].text
    assert "age: 30" in out.sections[0].blocks[0].text
    assert "name: bob" in out.sections[0].blocks[1].text


def test_extract_xlsx_unmerges_merged_cells(tmp_path: Path) -> None:
    """A MergedCell in a data row should resolve to its anchor value."""
    from openpyxl import Workbook

    xlsx = tmp_path / "merged.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Merged"
    # Headers are NOT in any merged range — keep them clean so column labels
    # are predictable.
    ws.append(["col1", "col2", "col3"])
    # Append a data row, THEN merge B2:C2 so the MergedCell falls inside
    # the data area only. The anchor (B2) keeps its original value and
    # C2 becomes a MergedCell that should resolve to the anchor on read.
    ws.append(["alpha", "beta", "gamma"])
    ws.merge_cells("B2:C2")
    wb.save(xlsx)
    wb.close()

    out = extract_xlsx(xlsx)
    cells = out.sections[0].blocks[0].text
    # col 1: its own header + own value
    assert "col1: alpha" in cells
    # col 2 (the merge anchor) — keeps "beta"
    assert "col2: beta" in cells
    # col 3 (MergedCell) — must resolve back to the anchor value "beta"
    assert "col3: beta" in cells


def test_extract_docx_groups_by_h1(tmp_path: Path) -> None:
    from docx import Document

    docx_path = tmp_path / "doc.docx"
    document = Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("alpha text")
    document.add_heading("Chapter Two", level=1)
    document.add_paragraph("beta text")
    document.save(docx_path)

    out = extract_docx(docx_path)
    assert out.format == "docx"
    assert [s.heading for s in out.sections] == ["Chapter One", "Chapter Two"]
    assert out.sections[0].blocks[0].text == "alpha text"
    assert out.sections[1].blocks[0].text == "beta text"


def test_extract_pptx_one_section_per_slide(tmp_path: Path) -> None:
    from pptx import Presentation

    pptx_path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide_layout = presentation.slide_layouts[1]  # title + content
    s1 = presentation.slides.add_slide(slide_layout)
    s1.shapes.title.text = "Intro"
    s1.placeholders[1].text = "first content"
    s2 = presentation.slides.add_slide(slide_layout)
    s2.shapes.title.text = "Body"
    s2.placeholders[1].text = "second content"
    presentation.save(pptx_path)

    out = extract_pptx(pptx_path)
    assert out.format == "pptx"
    assert [s.heading for s in out.sections] == ["Intro", "Body"]
    assert "first content" in out.sections[0].blocks[0].text


def test_extract_pdf_text_reads_embedded_layer(tmp_path: Path) -> None:
    import pymupdf

    pdf_path = tmp_path / "text.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello from a text-layer PDF.")
    doc.save(pdf_path)
    doc.close()

    out = extract_pdf_text(pdf_path)
    assert out.format == "pdf"
    assert out.page_count == 1
    assert "Hello from a text-layer PDF." in out.sections[0].blocks[0].text


def test_is_scanned_pdf_true_when_no_text_layer(tmp_path: Path) -> None:
    import pymupdf

    pdf_path = tmp_path / "blank.pdf"
    doc = pymupdf.open()
    doc.new_page()  # empty page, no text layer
    doc.save(pdf_path)
    doc.close()

    assert is_scanned_pdf(pdf_path, avg_chars_threshold=100) is True


def test_is_scanned_pdf_false_for_text_heavy_pdf(tmp_path: Path) -> None:
    import pymupdf

    pdf_path = tmp_path / "text_heavy.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), ("lorem ipsum " * 50))
    doc.save(pdf_path)
    doc.close()

    assert is_scanned_pdf(pdf_path, avg_chars_threshold=100) is False


# ---------------------------------------------------------------------------
# DeepDocClient — new /predict/* methods (HTTP transport mocked)
# ---------------------------------------------------------------------------


class _FakeTransport:
    """Minimal httpx transport that returns canned JSON responses."""

    def __init__(self, responses: list[dict[str, Any]]) -> None:
        self._responses = list(responses)
        self.calls: list[tuple[str, str, dict[str, str]]] = []

    def __call__(self, request):  # httpx signature
        from httpx import Response

        body_field = next(
            (k for k in request.headers if k.lower() == "content-type"), ""
        )
        self.calls.append((request.method, request.url.path, dict(request.headers)))
        idx = len(self.calls) - 1
        payload = self._responses[idx] if idx < len(self._responses) else {}
        return Response(200, json=payload, request=request)


def _client_with_transport(responses: list[dict[str, Any]]) -> DeepDocClient:
    """Build a DeepDocClient whose HTTP calls are answered by `_FakeTransport`."""
    import httpx

    transport = httpx.MockTransport(_FakeTransport(responses))
    http_client = httpx.Client(transport=transport, timeout=5.0)
    return DeepDocClient(base_url="http://deepdoc.test", api_key="ignored", client=http_client)


def test_deepdoc_client_health_uses_get_endpoint() -> None:
    client = _client_with_transport([{}])  # status 200, body ignored
    assert client.health() is True


def test_deepdoc_client_predict_dla_sends_jpeg_multipart() -> None:
    client = _client_with_transport([
        {"boxes": [{"type": "text", "bbox": [0, 0, 100, 50]}]},
    ])
    out = client.predict_dla(b"\xff\xd8\xff\xe0fake-jpeg", filename="p.jpg")
    assert out["boxes"][0]["type"] == "text"


def test_deepdoc_client_predict_ocr_passes_operator_data() -> None:
    client = _client_with_transport([{"text": "hello"}])
    out = client.predict_ocr(b"\xff\xd8\xff\xe0fake", operator="rec")
    assert out["text"] == "hello"


def test_deepdoc_client_predict_tsr_returns_html() -> None:
    client = _client_with_transport([{"html": "<table><tr><td>a</td></tr></table>"}])
    out = client.predict_tsr(b"\xff\xd8\xff\xe0fake")
    assert "<table>" in out["html"]


# ---------------------------------------------------------------------------
# DeepDocVisualRouter — orchestration of DLA + OCR + TSR into ParsedDoc
# ---------------------------------------------------------------------------


class _FakeDeepDocClient:
    """Records calls and returns canned DLA/OCR/TSR payloads.

    The shapes mirror the real DeepDoc LitServe server contract at port 9390
    — DLA ``bboxes`` are ``[x0, y0, x1, y1, score, class_id]`` with class
    5 = table; OCR-rec returns a 4-level nested ``output``; TSR returns a
    structural ``bboxes`` list (column / row boxes).  Visual router parsers
    in this codebase are tested against this exact shape.
    """

    def __init__(self) -> None:
        self.dla_calls = 0
        self.ocr_calls = 0
        self.tsr_calls = 0

    def predict_dla(self, image_bytes: bytes) -> dict[str, Any]:
        self.dla_calls += 1
        # class 1 = text, class 5 = table — mirrors the real DeepDoc convention
        return {
            "bboxes": [
                [10, 10, 200, 50, 0.95, 1.0],   # text
                [10, 100, 300, 250, 0.90, 5.0],  # table
            ]
        }

    def predict_ocr(self, image_bytes: bytes, *, operator: str = "rec") -> dict[str, Any]:
        self.ocr_calls += 1
        # 4-level nested: batch -> page -> items -> [text, score]
        return {"output": [[[["OCR line", 0.99]]]]}

    def predict_tsr(self, image_bytes: bytes) -> dict[str, Any]:
        self.tsr_calls += 1
        # Structural bboxes — class 1 = column, class 2 = row.
        return {
            "bboxes": [
                [10, 100, 300, 250, 0.90, 1.0],   # 1 column
                [10, 130, 300, 150, 0.85, 2.0],   # 1 row
            ]
        }


def _make_jpeg(width: int = 400, height: int = 300) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), color="white").save(buffer, format="JPEG")
    return buffer.getvalue()


def test_visual_router_aggregates_two_pages_into_two_sections() -> None:
    fake = _FakeDeepDocClient()
    router = DeepDocVisualRouter(client=fake, fallback_ocr=False)
    images = [_make_jpeg(), _make_jpeg()]
    out = router.parse_images(images, document_name="scanned.pdf")
    assert fake.dla_calls == 2
    # 2 pages * (1 text box + 1 table box) = 4 blocks total
    assert sum(len(s.blocks) for s in out.sections) == 4
    assert out.page_count == 2
    # First page text block comes from OCR
    page1 = out.sections[0]
    assert any(b.kind == "text" and b.text == "OCR line" for b in page1.blocks)
    # First page table block comes from TSR → per-cell OCR cascade.
    # The fake TSR returns 1 column × 1 row, so the per-cell path runs
    # one OCR-rec call (OCR-rec stub returns "OCR line") and produces
    # a real markdown table — not the legacy placeholder.
    assert any(
        b.kind == "table" and "OCR line" in b.text for b in page1.blocks
    )


def test_visual_router_falls_back_to_rapidocr_when_dla_fails() -> None:
    """If DLA raises, the whole batch should degrade to RapidOCR."""

    class _BoomClient:
        def predict_dla(self, image_bytes: bytes) -> dict[str, Any]:
            raise RuntimeError("deepdoc down")

    router = DeepDocVisualRouter(client=_BoomClient(), fallback_ocr=True)
    # Render text onto a JPEG so RapidOCR has something to read — a blank
    # image would trigger "RapidOCR produced no text" instead of exercising
    # the fallback path.
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (400, 100), color="white")
    ImageDraw.Draw(img).text((10, 40), "fallback text 12345", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    out = router.parse_images([buf.getvalue()], document_name="scan.pdf")
    assert out.ocr_used is True


def test_visual_router_raises_when_fallback_disabled_and_dla_fails() -> None:
    class _BoomClient:
        def predict_dla(self, image_bytes: bytes) -> dict[str, Any]:
            raise RuntimeError("deepdoc down")

    router = DeepDocVisualRouter(client=_BoomClient(), fallback_ocr=False)
    with pytest.raises(RuntimeError):
        router.parse_images([_make_jpeg()])


def test_visual_router_falls_back_to_rapidocr_when_ocr_rec_returns_empty() -> None:
    """OCR-rec returns empty text for tall multi-line crops (verified on
    scanned_pdf.pdf boxes [1]/[2]/[3] — all height ~125px, all empty).
    Visual router should retry the same crop with RapidOCR to recover the
    paragraph text rather than silently dropping the block."""

    class _EmptyOcrClient:
        """Returns a tall text box whose OCR-rec payload is empty."""

        def __init__(self) -> None:
            self.ocr_calls = 0

        def predict_dla(self, image_bytes: bytes) -> dict[str, Any]:
            return {"bboxes": [[10, 10, 400, 100, 0.9, 1.0]]}

        def predict_ocr(self, image_bytes: bytes, *, operator: str = "rec") -> dict[str, Any]:
            self.ocr_calls += 1
            # 4-level nested, but the items list is empty — what DeepDoc
            # OCR-rec actually returns for multi-line crops.
            return {"output": [[[]]]}

        def predict_tsr(self, image_bytes: bytes) -> dict[str, Any]:
            return {"bboxes": []}

    # Render text onto the JPEG so RapidOCR has something to read after
    # the OCR-rec empty-result branch fires.
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (400, 100), color="white")
    ImageDraw.Draw(img).text((10, 40), "RapidOCR fallback text", fill="black")
    buf = io.BytesIO()
    img.save(buf, format="JPEG")

    client = _EmptyOcrClient()
    router = DeepDocVisualRouter(client=client, fallback_ocr=True)
    out = router.parse_images([buf.getvalue()], document_name="scan.pdf")
    blocks = out.sections[0].blocks
    assert client.ocr_calls == 1, "OCR-rec should have been called once"
    # RapidOCR recovered the text from the same crop. OCR engines commonly
    # collapse whitespace, so check for a substring that survives spacing.
    assert any("RapidOCR" in b.text and "fallback" in b.text for b in blocks), (
        "RapidOCR fallback should have recovered the text; got blocks: "
        + repr([b.text for b in blocks])
    )


def test_visual_router_coerce_bbox_accepts_dict_form() -> None:
    bbox = DeepDocVisualRouter._coerce_bbox({"x0": 1, "y0": 2, "x1": 3, "y1": 4})
    assert bbox == BoundingBox(x0=1.0, y0=2.0, x1=3.0, y1=4.0)
    assert DeepDocVisualRouter._coerce_bbox("nope") is None


# ---------------------------------------------------------------------------
# parse_document_with_router — top-level dispatcher
# ---------------------------------------------------------------------------


def test_dispatcher_routes_md_to_plain_extractor(tmp_path: Path) -> None:
    md = tmp_path / "note.md"
    md.write_text("# Title\nbody text\n", encoding="utf-8")
    out = parse_document_with_router(md)
    assert out.format == "md"
    assert out.sections[0].heading == "Title"


def test_dispatcher_routes_txt_to_plain_extractor(tmp_path: Path) -> None:
    txt = tmp_path / "note.txt"
    txt.write_text("hello world\n" * 40, encoding="utf-8")
    out = parse_document_with_router(txt)
    assert out.format == "txt"
    assert out.sections  # non-empty


def test_dispatcher_routes_xlsx_to_openpyxl(tmp_path: Path) -> None:
    from openpyxl import Workbook

    xlsx = tmp_path / "s.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["k", "v"])
    ws.append(["foo", "bar"])
    wb.save(xlsx)
    wb.close()
    out = parse_document_with_router(xlsx)
    assert out.format == "xlsx"
    assert "k: foo" in out.sections[0].blocks[0].text


def test_dispatcher_routes_text_pdf_to_pymupdf(tmp_path: Path) -> None:
    import pymupdf

    pdf = tmp_path / "t.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "embedded text")
    doc.save(pdf)
    doc.close()
    # No visual_router passed -> scanner check is bypassed, falls back to PyMuPDF.
    out = parse_document_with_router(pdf)
    assert out.format == "pdf"
    assert "embedded text" in out.sections[0].blocks[0].text


def test_dispatcher_routes_scanned_pdf_to_visual_router(tmp_path: Path) -> None:
    import pymupdf

    pdf = tmp_path / "blank.pdf"
    doc = pymupdf.open()
    doc.new_page()
    doc.save(pdf)
    doc.close()

    fake = _FakeDeepDocClient()
    router = DeepDocVisualRouter(client=fake, fallback_ocr=False)
    out = parse_document_with_router(pdf, visual_router=router)
    assert fake.dla_calls == 1
    assert out.page_count == 1


def test_dispatcher_rejects_unknown_extension(tmp_path: Path) -> None:
    bad = tmp_path / "evil.exe"
    bad.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_document_with_router(bad)


# =============================================================================
# DLA class 0/1/2/3/4/5/6/8 full-coverage tests (T4.4 visual_router)
# =============================================================================


class _ClassRoutedFakeClient:
    """Fake DeepDocClient that returns a caller-supplied list of DLA bboxes.

    Each bbox is ``[x0, y0, x1, y1, score, class_id]``. OCR-rec returns a
    stub string; TSR returns caller-supplied bboxes (used to drive the
    per-cell OCR cascade). Used by the per-class routing tests below to
    exercise one class_id at a time without depending on the legacy
    two-class fake.
    """

    def __init__(
        self,
        dla_bboxes: list[list[float]],
        tsr_bboxes: list[list[float]] | None = None,
        ocr_text: str = "OCR line",
        ocr_rec_payloads: list[dict[str, Any]] | None = None,
    ) -> None:
        self._dla_bboxes = dla_bboxes
        self._tsr_bboxes = tsr_bboxes or []
        self._ocr_text = ocr_text
        self._ocr_rec_payloads = list(ocr_rec_payloads or [])
        self.dla_calls = 0
        self.ocr_calls = 0
        self.tsr_calls = 0

    def predict_dla(self, image_bytes: bytes) -> dict[str, Any]:
        self.dla_calls += 1
        return {"bboxes": list(self._dla_bboxes)}

    def predict_ocr(self, image_bytes: bytes, *, operator: str = "rec") -> dict[str, Any]:
        self.ocr_calls += 1
        if self._ocr_rec_payloads:
            return self._ocr_rec_payloads.pop(0)
        return {"output": [[[[self._ocr_text, 0.99]]]]}

    def predict_tsr(self, image_bytes: bytes) -> dict[str, Any]:
        self.tsr_calls += 1
        return {"bboxes": list(self._tsr_bboxes)}


def _parse_with_client(fake: _ClassRoutedFakeClient) -> list:
    """Run one image through the visual router and return all blocks."""
    router = DeepDocVisualRouter(client=fake, fallback_ocr=False)
    out = router.parse_images([_make_jpeg()], document_name="x.pdf")
    return [b for s in out.sections for b in s.blocks]


def test_class_2_reference_routes_as_text() -> None:
    """Class-2 (reference list) used to be dropped — now treated as text."""
    fake = _ClassRoutedFakeClient(dla_bboxes=[[10, 10, 400, 100, 0.9, 2.0]])
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    assert blocks[0].kind == "text"
    assert blocks[0].text == "OCR line"
    assert fake.ocr_calls == 1


def test_class_3_figure_routes_as_bbox_placeholder() -> None:
    """Class-3 (figure) never goes through OCR — bbox-only placeholder."""
    fake = _ClassRoutedFakeClient(dla_bboxes=[[10, 10, 400, 100, 0.9, 3.0]])
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    block = blocks[0]
    assert block.kind == "figure"
    assert "figure @ page 1" in block.text
    assert "bbox=" in block.text
    # No OCR / TSR call — figure is just a placeholder.
    assert fake.ocr_calls == 0
    assert fake.tsr_calls == 0
    # bbox metadata preserved on the block for downstream consumers.
    assert block.meta.get("dla_bbox") is not None
    assert block.meta["dla_bbox"]["x0"] == 10.0


def test_class_4_figure_caption_routes_as_atomic_text() -> None:
    """Class-4 (figure caption) gets OCR-rec + atomic kind in chunker."""
    fake = _ClassRoutedFakeClient(dla_bboxes=[[10, 10, 400, 100, 0.9, 4.0]])
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    assert blocks[0].kind == "figure_caption"
    assert blocks[0].text == "OCR line"


def test_class_6_table_caption_routes_as_atomic_text() -> None:
    """Class-6 (table caption) — separate kind, OCR-rec text."""
    fake = _ClassRoutedFakeClient(dla_bboxes=[[10, 10, 400, 100, 0.9, 6.0]])
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    assert blocks[0].kind == "table_caption"
    assert blocks[0].text == "OCR line"


def test_class_8_equation_routes_as_atomic_text() -> None:
    """Class-8 (equation) — OCR-rec text, atomic in chunker."""
    fake = _ClassRoutedFakeClient(dla_bboxes=[[10, 10, 400, 100, 0.9, 8.0]])
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    assert blocks[0].kind == "equation"
    assert blocks[0].text == "OCR line"


def test_class_5_table_html_fast_path_used_when_present() -> None:
    """When TSR returns ``html``, parse it as markdown instead of per-cell."""
    html = (
        "<table>"
        "<tr><th>A</th><th>B</th></tr>"
        "<tr><td>1</td><td>2</td></tr>"
        "</table>"
    )
    fake = _ClassRoutedFakeClient(
        dla_bboxes=[[10, 100, 300, 250, 0.9, 5.0]],
        tsr_bboxes=[],  # no structural bboxes — would force placeholder
    )
    # Patch the payload to include html at predict_tsr time.
    def _tsr_with_html(_image: bytes) -> dict[str, Any]:
        fake.tsr_calls += 1
        return {"html": html, "bboxes": []}
    fake.predict_tsr = _tsr_with_html  # type: ignore[assignment]
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    text = blocks[0].text
    # Fast-path must not invoke per-cell OCR.
    assert fake.ocr_calls == 0
    assert "| A | B |" in text
    assert "| --- | --- |" in text
    assert "| 1 | 2 |" in text


def test_format_tsr_as_html_handles_attributes() -> None:
    """Lock down the HTML fast-path parser against non-trivial HTML.

    The live DeepDoc server doesn't return ``html`` today (verified 18/18
    tables on 2026-08-24), so this method is a future-proof hook. We
    still need to know it handles the kind of HTML DeepDoc might
    realistically emit next time it lights up: attribute-bearing tags,
    spanning cells, self-closing ``<br/>``, and uppercase tag names.

    Each case calls the static method directly so the test stays a
    pure-HTML contract, independent of the visual_router state machine.
    """
    fmt = DeepDocVisualRouter._format_tsr_as_html

    # Case 1: ``<table border="1">`` style attributes on every tag.
    md1 = fmt({
        "html": (
            '<table border="1" class="data">'
            '<tr><th class="h">A</th><th>B</th></tr>'
            '<tr><td rowspan="2" colspan="1">1</td><td>2</td></tr>'
            '</table>'
        )
    })
    assert md1 is not None
    assert "| A |" in md1 and "| B |" in md1
    assert "| --- | --- |" in md1
    assert "| 1 | 2 |" in md1
    # Attributes must not bleed into cell text.
    assert "border" not in md1
    assert "rowspan" not in md1 and "colspan" not in md1

    # Case 2: self-closing <br/> inside a cell (DeepDoc historical output).
    # The stdlib HTMLParser does not synthesize a newline for <br/>, so the
    # text after the tag surfaces into the same cell with a leading space.
    # The contract we lock down here is "does not crash + text before
    # <br/> is preserved + table structure stays intact". If a future
    # DeepDoc release leans on <br/> for multiline cells, this test will
    # need a parser upgrade; in the meantime the empty-content stop in
    # handle_data() means post-<br/> text can replace pre-<br/> text in
    # the same cell — that's the current behaviour and it's documented.
    md2 = fmt({
        "html": (
            "<table>"
            "<tr><td>line one<br/>line two</td><td>ok</td></tr>"
            "<tr><td>next</td><td>row</td></tr>"
            "</table>"
        )
    })
    assert md2 is not None
    flat2 = " ".join(md2.split())
    # The pre-<br/> text OR post-<br/> text must survive (not both, given
    # the empty-content gate), and the table's structural shape stays.
    assert ("line one" in flat2 or "line two" in flat2)
    assert "| ok |" in md2
    assert "| next | row |" in md2
    assert "| --- | --- |" in md2

    # Case 3: uppercase tag names (``<TR>``, ``<TD>``).
    md3 = fmt({
        "html": (
            "<TABLE>"
            "<TR><TH>X</TH></TR>"
            "<TR><TD>1</TD></TR>"
            "</TABLE>"
        )
    })
    assert md3 is not None
    assert "| X |" in md3
    assert "| 1 |" in md3

    # Case 4: empty / missing html -> None (caller falls back to per-cell OCR).
    assert fmt({"html": ""}) is None
    assert fmt({"html": None}) is None
    assert fmt({}) is None

    # Case 5: html present but no <table> -> None (parser rows empty).
    assert fmt({"html": "<div>not a table</div>"}) is None


def test_class_5_table_per_cell_ocr_assembles_markdown() -> None:
    """TSR returns 2 rows × 3 cols → per-cell OCR called 6×, md assembled."""
    tsr_bboxes = [
        # 3 columns
        [10, 100, 100, 250, 0.9, 1.0],
        [100, 100, 200, 250, 0.9, 1.0],
        [200, 100, 300, 250, 0.9, 1.0],
        # 2 rows
        [10, 100, 300, 175, 0.9, 2.0],
        [10, 175, 300, 250, 0.9, 2.0],
    ]
    fake = _ClassRoutedFakeClient(
        dla_bboxes=[[10, 100, 300, 250, 0.9, 5.0]],
        tsr_bboxes=tsr_bboxes,
        ocr_text="cell",
    )
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    assert blocks[0].kind == "table"
    # 6 cells → 6 OCR-rec calls. (RapidOCR fallback is not exercised
    # because the fake returns non-empty text.)
    assert fake.ocr_calls == 6
    text = blocks[0].text
    # Emitter writes a 1-row header (empty) then 2 body rows of 3 cells;
    # count "cell" occurrences instead of full pipe-bracketed cells.
    assert text.count("cell") == 6
    assert "| --- | --- | --- |" in text


def test_class_5_table_budget_exceeded_emits_placeholder() -> None:
    """A 50×50 grid (2500 cells) must fall back to the placeholder text."""
    # 50 columns + 50 rows.
    tsr_bboxes: list[list[float]] = []
    for i in range(50):
        tsr_bboxes.append([i * 5, 100, i * 5 + 5, 250, 0.9, 1.0])
    for j in range(50):
        tsr_bboxes.append([10, j * 5, 250, j * 5 + 5, 0.9, 2.0])
    fake = _ClassRoutedFakeClient(
        dla_bboxes=[[10, 100, 300, 250, 0.9, 5.0]],
        tsr_bboxes=tsr_bboxes,
    )
    blocks = _parse_with_client(fake)
    assert len(blocks) == 1
    assert "DeepDoc table" in blocks[0].text
    assert "50 columns × 50 rows" in blocks[0].text
    # No per-cell OCR calls — budget gate fires before any crop.
    assert fake.ocr_calls == 0


def test_unknown_class_id_is_dropped_with_warning(caplog) -> None:
    """Unknown class_id (e.g. 7) must be dropped, not silently passed.

    When every box in a page is dropped, ``parse_images`` raises
    ``RuntimeError("deepdoc returned no usable sections")`` — this is
    correct: with nothing to render, the document is unusable. The
    WARNING is logged before the RuntimeError, so caplog captures it
    regardless of which exception bubbles up.
    """
    import logging

    import pytest

    fake = _ClassRoutedFakeClient(dla_bboxes=[[10, 10, 400, 100, 0.9, 7.0]])
    with caplog.at_level(
        logging.WARNING,
        logger="agentic_rag_project.doc_processor.visual_router",
    ):
        with pytest.raises(RuntimeError, match="no usable sections"):
            _parse_with_client(fake)
    assert any("unknown DLA class_id 7.0" in r.message for r in caplog.records)
